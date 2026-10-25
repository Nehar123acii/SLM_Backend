import csv
import os
import re
import openpyxl
import pytest
from sqlalchemy.orm import Session

from backend.db.models import Column, Dataset, Occurrence, Person
from backend.security import hmac_value
from backend.services.inventory import InventoryService
from backend.services.provenance import (
    rows_for_field,
    people_for_dataset,
    sources_for_person,
)

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), "..", "sample_data")


def test_independent_counts_patients_xlsx(db_session: Session):
    """
    Verifies occurrences per column against an independent count
    made directly with openpyxl on patients.xlsx.
    """
    file_path = os.path.join(SAMPLE_DIR, "patients.xlsx")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_idx = {name: idx + 1 for idx, name in enumerate(headers)}

    expected_counts = {}
    for col_name in ["Name", "Phone", "Email", "Aadhaar"]:
        c_idx = col_idx[col_name]
        cnt = 0
        for row in range(2, ws.max_row + 1):
            val = ws.cell(row=row, column=c_idx).value
            if val is not None and str(val).strip() != "":
                cnt += 1
        expected_counts[col_name.lower()] = cnt

    with open(file_path, "rb") as f:
        bytes_data = f.read()

    res = InventoryService.process_upload(
        db=db_session,
        original_filename="patients.xlsx",
        file_bytes=bytes_data,
        source_type="excel",
    )
    dataset_id = res["dataset_id"]

    for col_name, expected in expected_counts.items():
        col = (
            db_session.query(Column)
            .filter_by(dataset_id=dataset_id, column_name=col_name)
            .first()
        )
        assert col is not None, f"Column {col_name} not found"
        assert col.locator is not None and "patients.xlsx" in col.locator
        assert col.category in ["name", "phone", "email", "aadhaar"]
        assert col.impact_tier in ["2", "3"]

        actual_occurrences = (
            db_session.query(Occurrence).filter_by(column_id=col.id).count()
        )
        assert actual_occurrences == expected, (
            f"Occurrence mismatch for {col_name}: got {actual_occurrences}, expected {expected}"
        )


def test_independent_counts_billing_csv(db_session: Session):
    """
    Verifies occurrences per column against an independent count
    made directly with python's standard csv module on billing.csv.
    """
    file_path = os.path.join(SAMPLE_DIR, "billing.csv")
    expected_counts = {"patient_name": 0, "phone": 0}

    with open(file_path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            for k in expected_counts:
                if row.get(k) and row[k].strip():
                    expected_counts[k] += 1

    with open(file_path, "rb") as f:
        bytes_data = f.read()

    res = InventoryService.process_upload(
        db=db_session,
        original_filename="billing.csv",
        file_bytes=bytes_data,
        source_type="csv",
    )
    dataset_id = res["dataset_id"]

    for col_name, expected in expected_counts.items():
        col = (
            db_session.query(Column)
            .filter_by(dataset_id=dataset_id, column_name=col_name)
            .first()
        )
        assert col is not None
        assert col.locator == f"billing.csv||{col_name}"
        actual = db_session.query(Occurrence).filter_by(column_id=col.id).count()
        assert actual == expected


def test_independent_counts_appointments_form_csv(db_session: Session):
    """
    Verifies occurrences per column against an independent count
    made directly with csv module on appointments_form.csv.
    """
    file_path = os.path.join(SAMPLE_DIR, "appointments_form.csv")
    expected_counts = {"name": 0, "phone": 0, "email": 0}

    with open(file_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get("Name") and row["Name"].strip():
                expected_counts["name"] += 1
            if row.get("Phone") and row["Phone"].strip():
                expected_counts["phone"] += 1
            if row.get("Email") and row["Email"].strip():
                expected_counts["email"] += 1

    with open(file_path, "rb") as f:
        bytes_data = f.read()

    res = InventoryService.process_upload(
        db=db_session,
        original_filename="appointments_form.csv",
        file_bytes=bytes_data,
        source_type="google_forms",
    )
    dataset_id = res["dataset_id"]

    for col_name, expected in expected_counts.items():
        col = (
            db_session.query(Column)
            .filter_by(dataset_id=dataset_id, column_name=col_name)
            .first()
        )
        assert col is not None
        actual = db_session.query(Occurrence).filter_by(column_id=col.id).count()
        assert actual == expected


def test_independent_counts_staff_csv(db_session: Session):
    """
    Verifies occurrences per column against an independent count
    made directly with csv module on staff.csv.
    """
    file_path = os.path.join(SAMPLE_DIR, "staff.csv")
    expected_counts = {"name": 0, "pan": 0, "phone": 0}

    with open(file_path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            for k in expected_counts:
                if row.get(k) and row[k].strip():
                    expected_counts[k] += 1

    with open(file_path, "rb") as f:
        bytes_data = f.read()

    res = InventoryService.process_upload(
        db=db_session,
        original_filename="staff.csv",
        file_bytes=bytes_data,
        source_type="csv",
    )
    dataset_id = res["dataset_id"]

    for col_name, expected in expected_counts.items():
        col = (
            db_session.query(Column)
            .filter_by(dataset_id=dataset_id, column_name=col_name)
            .first()
        )
        assert col is not None
        actual = db_session.query(Occurrence).filter_by(column_id=col.id).count()
        assert actual == expected


def test_independent_counts_whatsapp_chat(db_session: Session):
    """
    Verifies occurrences in sender column against an independent count
    made directly with line-parsing regex on whatsapp_chat.txt.
    """
    file_path = os.path.join(SAMPLE_DIR, "whatsapp_chat.txt")
    re_msg = re.compile(r"^\d{2}/\d{2}/\d{2,4},.*?\s-\s([^:]+):")

    expected_sender_count = 0
    with open(file_path, "r", encoding="utf-8") as f:
        for line in f:
            m = re_msg.match(line)
            if m:
                sender = m.group(1).strip()
                if sender.upper() not in ["SYSTEM", "MESSAGES AND CALLS"]:
                    expected_sender_count += 1

    with open(file_path, "rb") as f:
        bytes_data = f.read()

    res = InventoryService.process_upload(
        db=db_session,
        original_filename="whatsapp_chat.txt",
        file_bytes=bytes_data,
        source_type="whatsapp",
    )
    dataset_id = res["dataset_id"]

    sender_col = (
        db_session.query(Column)
        .filter_by(dataset_id=dataset_id, column_name="sender")
        .first()
    )
    assert sender_col is not None
    assert sender_col.category == "name"

    actual_senders = (
        db_session.query(Occurrence).filter_by(column_id=sender_col.id).count()
    )
    assert actual_senders == expected_sender_count


def test_person_appearing_in_patients_and_billing_gets_one_person_id(db_session: Session):
    """
    Verifies strong match: a person appearing in patients.xlsx and billing.csv
    with the same phone in different formats gets ONE person_id.
    """
    with open(os.path.join(SAMPLE_DIR, "patients.xlsx"), "rb") as f:
        res1 = InventoryService.process_upload(
            db=db_session,
            original_filename="patients.xlsx",
            file_bytes=f.read(),
            source_type="excel",
        )

    with open(os.path.join(SAMPLE_DIR, "billing.csv"), "rb") as f:
        res2 = InventoryService.process_upload(
            db=db_session,
            original_filename="billing.csv",
            file_bytes=f.read(),
            source_type="csv",
        )

    # Find the phone column in both datasets
    col_pat_phone = (
        db_session.query(Column)
        .filter_by(dataset_id=res1["dataset_id"], column_name="phone")
        .first()
    )
    col_bil_phone = (
        db_session.query(Column)
        .filter_by(dataset_id=res2["dataset_id"], column_name="phone")
        .first()
    )

    assert col_pat_phone is not None
    assert col_bil_phone is not None

    # Get occurrences for patients and billing
    pat_occs = db_session.query(Occurrence).filter_by(column_id=col_pat_phone.id).all()
    bil_occs = db_session.query(Occurrence).filter_by(column_id=col_bil_phone.id).all()

    # Find an overlapping phone HMAC
    pat_hmac_map = {o.value_hmac: o for o in pat_occs}
    overlapping_found = 0

    for b_occ in bil_occs:
        if b_occ.value_hmac in pat_hmac_map:
            p_occ = pat_hmac_map[b_occ.value_hmac]
            assert b_occ.person_id == p_occ.person_id, (
                f"Expected same person_id for matching phone HMAC, got {b_occ.person_id} vs {p_occ.person_id}"
            )
            person = db_session.get(Person, b_occ.person_id)
            assert person.resolution == "strong"
            assert person.confidence == 1.0
            overlapping_found += 1

    assert overlapping_found > 0, "No overlapping phones found between patients and billing"


def test_name_only_match_stays_probable_and_not_merged(db_session: Session):
    """
    Verifies that records matching on name only remain distinct probable persons
    and are never merged into one person_id.
    """
    csv1 = (
        "patient_name,phone\n"
        "Same Person Name,+91 91111 22222\n"
    )
    csv2 = (
        "patient_name,phone\n"
        "Same Person Name,+91 98888 77777\n"
    )

    res1 = InventoryService.process_upload(
        db=db_session,
        original_filename="name_test_1.csv",
        file_bytes=csv1.encode("utf-8"),
        source_type="csv",
    )
    res2 = InventoryService.process_upload(
        db=db_session,
        original_filename="name_test_2.csv",
        file_bytes=csv2.encode("utf-8"),
        source_type="csv",
    )

    col1 = db_session.query(Column).filter_by(dataset_id=res1["dataset_id"], column_name="patient_name").first()
    col2 = db_session.query(Column).filter_by(dataset_id=res2["dataset_id"], column_name="patient_name").first()

    occ1 = db_session.query(Occurrence).filter_by(column_id=col1.id).first()
    occ2 = db_session.query(Occurrence).filter_by(column_id=col2.id).first()

    assert occ1 is not None and occ2 is not None
    # Because phones are different, they MUST NOT be merged
    assert occ1.person_id != occ2.person_id

    # Test pure name-only upload (no phone or email)
    csv_names_only = (
        "name,notes\n"
        "Unverified Principal,walk in inquiry\n"
    )
    res3 = InventoryService.process_upload(
        db=db_session,
        original_filename="name_only.csv",
        file_bytes=csv_names_only.encode("utf-8"),
        source_type="csv",
    )
    col3 = db_session.query(Column).filter_by(dataset_id=res3["dataset_id"], column_name="name").first()
    occ3 = db_session.query(Occurrence).filter_by(column_id=col3.id).first()
    assert occ3 is not None
    person3 = db_session.get(Person, occ3.person_id)
    assert person3.resolution == "probable"
    assert person3.confidence == 0.6


def test_low_confidence_ocr_does_not_create_strong_link(db_session: Session):
    """
    Verifies that OCR occurrences with ocr_conf below 0.8 cannot create a strong link.
    """
    import io
    from unittest.mock import patch
    from PIL import Image
    import pandas as pd
    from backend.etl.ocr_engine import DocumentExtractor

    # 1. First upload a verified rule-based patient
    rule_csv = (
        "name,phone\n"
        "Dr. Sharma,+91 99999 88888\n"
    )
    res_rule = InventoryService.process_upload(
        db=db_session,
        original_filename="rule_patient.csv",
        file_bytes=rule_csv.encode("utf-8"),
        source_type="csv",
    )
    col_rule = db_session.query(Column).filter_by(dataset_id=res_rule["dataset_id"], column_name="phone").first()
    occ_rule = db_session.query(Occurrence).filter_by(column_id=col_rule.id).first()
    person_rule_id = occ_rule.person_id

    # 2. Now upload a valid PNG document containing the same phone but low confidence (0.5 < 0.8)
    img = Image.new("RGB", (300, 100), color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    png_bytes = buf.getvalue()

    mock_df = pd.DataFrame([{
        "record_ref": "rec_1",
        "page_number": 1,
        "line_number": 1,
        "text_content": "Contact: +91 99999 88888",
        "extraction_method": "tesseract",
        "confidence": 0.50,
        "char_length": 25,
        "status": "CONFIDENT",
        "source_reference": "scan.png",
    }])

    with patch.object(DocumentExtractor, "extract", return_value=(mock_df, [], {"text_content": "text_content"}, {"source_type": "document_ocr"})):
        res_ocr = InventoryService.process_upload(
            db=db_session,
            original_filename="scan.png",
            file_bytes=png_bytes,
            source_type="document_ocr",
        )

    col_ocr = db_session.query(Column).filter_by(dataset_id=res_ocr["dataset_id"], column_name="text_content").first()
    occ_ocr = db_session.query(Occurrence).filter_by(column_id=col_ocr.id).first()

    assert occ_ocr is not None
    assert occ_ocr.origin == "ocr"
    assert occ_ocr.ocr_conf == 0.50
    # Must NOT have merged into person_rule_id because ocr_conf < 0.8
    assert occ_ocr.person_id != person_rule_id


def test_provenance_service_helpers(db_session: Session):
    """
    Verifies rows_for_field, people_for_dataset, and sources_for_person helpers.
    """
    csv_data = (
        "patient_name,phone,email\n"
        "Aditi Rao,+91 98765 11111,aditi@example.com\n"
    )
    res = InventoryService.process_upload(
        db=db_session,
        original_filename="helper_test.csv",
        file_bytes=csv_data.encode("utf-8"),
        source_type="csv",
    )
    dataset_id = res["dataset_id"]

    col_phone = db_session.query(Column).filter_by(dataset_id=dataset_id, column_name="phone").first()
    assert col_phone is not None

    # Test rows_for_field
    field_rows = rows_for_field(col_phone.id, db=db_session)
    assert len(field_rows) == 1
    assert field_rows[0]["masked_sample"] == "XXXXXX1111"
    assert "98765" not in field_rows[0]["masked_sample"]

    # Test people_for_dataset
    people = people_for_dataset(dataset_id, db=db_session)
    assert len(people) == 1
    pid = people[0].id

    # Test sources_for_person
    sources = sources_for_person(pid, db=db_session)
    assert len(sources) >= 1
    col_names = [s["column_name"] for s in sources]
    assert "phone" in col_names
    assert "helper_test.csv" in sources[0]["locator"]
